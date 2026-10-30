
import io
import re
import zipfile
import logging

import pdfplumber
from docx import Document


logger = logging.getLogger(__name__)

MAX_FILE_SIZE = 5 * 1024 * 1024  # 5 MB


class ResumeParsingError(Exception):
    """Raised when a resume cannot be safely or correctly parsed."""
    pass


def extract_resume_text(filename: str, file_bytes: bytes) -> str:
    """
    Extract text from PDF or DOCX resume files.

    Args:
        filename: Original uploaded filename.
        file_bytes: Uploaded file contents.

    Returns:
        Extracted resume text.

    Raises:
        ResumeParsingError: If the file is invalid, unsupported,
        too large, unreadable, or contains no extractable text.
    """

    if not filename or not file_bytes:
        raise ResumeParsingError("Please upload a valid resume.")

    if len(file_bytes) > MAX_FILE_SIZE:
        raise ResumeParsingError("Resume exceeds the 5 MB size limit.")

    extension = filename.lower().rsplit(".", 1)[-1]

    if extension not in {"pdf", "docx"}:
        raise ResumeParsingError(
            "Unsupported file format. Please upload PDF or DOCX."
        )

    try:
        if extension == "pdf":
            if not file_bytes.startswith(b"%PDF-"):
                raise ResumeParsingError("Invalid PDF file.")

            text = extract_pdf_text(file_bytes)

        else:
            if not zipfile.is_zipfile(io.BytesIO(file_bytes)):
                raise ResumeParsingError("Invalid DOCX file.")

            text = extract_docx_text(file_bytes)

        # Normalise whitespace without changing the original wording.
        text = re.sub(r"\s+", " ", text).strip()

        if len(text) < 20:
            raise ResumeParsingError(
                "Could not extract enough text from this resume. "
                "If it is a scanned PDF, OCR support is required."
            )

        return text

    except ResumeParsingError:
        raise

    except Exception as exc:
        logger.exception("Resume parsing failed")
        raise ResumeParsingError(
            "Unable to read this resume. Please check the file."
        ) from exc


def extract_pdf_text(file_bytes: bytes) -> str:
    """Extract text from all pages of a PDF."""

    extracted_pages = []

    with pdfplumber.open(io.BytesIO(file_bytes)) as pdf:
        if not pdf.pages:
            raise ResumeParsingError("The PDF contains no pages.")

        for page in pdf.pages:
            page_text = page.extract_text()

            if page_text:
                extracted_pages.append(page_text)

    return "\n".join(extracted_pages)


def extract_docx_text(file_bytes: bytes) -> str:
    """Extract text from DOCX paragraphs and tables."""

    document = Document(io.BytesIO(file_bytes))
    extracted_text = []

    # Extract normal paragraphs.
    for paragraph in document.paragraphs:
        if paragraph.text.strip():
            extracted_text.append(paragraph.text)

    # Extract tables, which are common in resumes.
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    extracted_text.append(cell.text)

    return "\n".join(extracted_text)
